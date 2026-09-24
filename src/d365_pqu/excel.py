from __future__ import annotations

import json
from collections.abc import Callable, Iterable, Mapping, Sequence
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from d365_pqu.config import DATASET_NAME, SCHEMA_VERSION
from d365_pqu.models import NormalizedDataset
from d365_pqu.serialization import isoformat

TITLE_FILL = PatternFill("solid", fgColor="0B3D5C")
HEADER_FILL = PatternFill("solid", fgColor="12506E")
SECTION_FILL = PatternFill("solid", fgColor="E8EEF3")
WHITE_BOLD = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16)
LABEL_FONT = Font(bold=True)
LINK_FONT = Font(color="0B5FA5", underline="single")
THIN = Side(style="thin", color="C9D4DD")
BORDER = Border(bottom=THIN)
STATUS_FILLS = {
    "Completed": PatternFill("solid", fgColor="DDF3E4"),
    "In-Progress": PatternFill("solid", fgColor="FFF3CD"),
    "Not Started": PatternFill("solid", fgColor="E7ECF2"),
    "Canceled": PatternFill("solid", fgColor="F8D7DA"),
}
DATE_FORMAT = "yyyy-mm-dd"
DATE_FIELDS = {
    "change_cutoff_date",
    "train_start_date",
    "train_end_date",
    "sandbox_start_date",
    "sandbox_end_date",
    "production_start_date",
    "production_end_date",
}

Column = tuple[str, str | Callable[[Mapping[str, Any]], Any], int]


def _value(record: Mapping[str, Any], key: str) -> Any:
    if key.startswith("source."):
        source = record.get("source") or {}
        return source.get(key.split(".", 1)[1])
    return record.get(key)


def _excel_value(value: Any, key: str) -> Any:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)
    if isinstance(value, str):
        if key in DATE_FIELDS:
            try:
                return date.fromisoformat(value)
            except ValueError:
                pass
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _write_table_sheet(
    wb: WorkbookType,
    title: str,
    columns: Sequence[Column],
    records: Iterable[Mapping[str, Any]],
    *,
    freeze: str = "A2",
    table_name: str | None = None,
) -> Worksheet:
    ws = wb.create_sheet(title)
    headers = [column[0] for column in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    rows = list(records)
    for record in rows:
        values: list[Any] = []
        for _, key, _ in columns:
            value = key(record) if callable(key) else _value(record, key)
            field_name = key if isinstance(key, str) else ""
            values.append(_excel_value(value, field_name))
        ws.append(values)
    for index, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(columns)):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top")
    status_index = _status_column(columns)
    if status_index:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=status_index)
            fill = STATUS_FILLS.get(str(cell.value))
            if fill:
                cell.fill = fill
    if rows and table_name:
        reference = f"A1:{get_column_letter(len(columns))}{ws.max_row}"
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight9",
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws


def _status_column(columns: Sequence[Column]) -> int | None:
    for index, (header, _, _) in enumerate(columns, start=1):
        if header == "Status":
            return index
    return None


def _readme(
    wb: WorkbookType,
    dataset: NormalizedDataset,
    metadata: dict[str, Any],
    health: dict[str, Any],
) -> Worksheet:
    ws = wb.create_sheet("00_README", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 92
    ws.merge_cells("A1:B1")
    title = ws["A1"]
    title.value = "D365 Finance & Operations PQU Tracker"
    title.font = TITLE_FONT
    title.fill = TITLE_FILL
    title.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 30
    rows: list[tuple[str, Any]] = [
        ("Dataset", DATASET_NAME),
        ("Schema version", SCHEMA_VERSION),
        ("Last synchronized", metadata["generated_at"]),
        ("Latest application version", dataset.records[-1]["application_version"]),
        ("Latest active PQU", metadata.get("latest_pqu_id") or "None"),
        ("Active PQU trains", metadata["current_count"]),
        ("Upcoming PQU trains", metadata["upcoming_count"]),
        ("Station schedule rows", metadata["station_schedule_count"]),
        ("Region mappings", metadata["region_count"]),
        ("Source", "Microsoft Learn public documentation"),
        ("Source article", dataset.source["article_url"]),
        ("Source commit", dataset.source["commit"]),
        ("Source SHA-256", dataset.source["sha256"]),
        ("Publication status", health["status"]),
        ("Public dataset", metadata["links"]["pages"]),
        ("JSON API", metadata["links"]["api_pqu"]),
        ("CSV API", metadata["links"]["api_csv"]),
        ("Workbook download", metadata["links"]["workbook"]),
    ]
    current_row = 3
    for label, value in rows:
        ws.cell(row=current_row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=current_row, column=1).fill = SECTION_FILL
        cell = ws.cell(row=current_row, column=2, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if label in {
            "Source article",
            "Public dataset",
            "JSON API",
            "CSV API",
            "Workbook download",
        }:
            cell.hyperlink = str(value)
            cell.font = LINK_FONT
        current_row += 1
    note_row = current_row + 1
    ws.cell(
        row=note_row,
        column=1,
        value=(
            "This workbook is generated automatically from Microsoft's public PQU schedule. "
            "Detailed schedules are published shortly before a train starts and can change. "
            "Environment-specific update timing depends on Microsoft notifications."
        ),
    ).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=2)
    chart_start = current_row + 4
    ws.cell(row=chart_start, column=1, value="Status").font = LABEL_FONT
    ws.cell(row=chart_start, column=2, value="Count").font = LABEL_FONT
    counts = metadata.get("status_counts", {})
    for offset, (status, count) in enumerate(sorted(counts.items()), start=1):
        ws.cell(row=chart_start + offset, column=1, value=status)
        ws.cell(row=chart_start + offset, column=2, value=count)
    chart = BarChart()
    chart.title = "PQU status distribution"
    chart.height = 7
    chart.width = 16
    data = Reference(
        ws,
        min_col=2,
        min_row=chart_start,
        max_row=chart_start + len(counts),
    )
    categories = Reference(
        ws,
        min_col=1,
        min_row=chart_start + 1,
        max_row=chart_start + len(counts),
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, f"D{chart_start}")
    return ws


def _key_value_sheet(
    wb: WorkbookType,
    title: str,
    rows: Sequence[tuple[str, Any]],
    *,
    widths: tuple[int, int] = (34, 100),
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = widths[0]
    ws.column_dimensions["B"].width = widths[1]
    for row_index, (label, value) in enumerate(rows, start=1):
        label_cell = ws.cell(row=row_index, column=1, value=label)
        label_cell.font = LABEL_FONT
        label_cell.fill = SECTION_FILL
        value_cell = ws.cell(row=row_index, column=2, value=value)
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def _status_history_rows(
    dataset: NormalizedDataset, generated_at: datetime
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for change in dataset.changes:
        if change["field"] != "status":
            continue
        rows.append(
            {
                "changed_at": change["changed_at"],
                "pqu_id": change["pqu_id"],
                "change_type": change["change_type"],
                "previous_status": change["old_value"],
                "new_status": change["new_value"],
                "source_commit": change["source_commit"],
            }
        )
    if not rows:
        baseline = isoformat(generated_at)
        for record in dataset.records:
            rows.append(
                {
                    "changed_at": baseline,
                    "pqu_id": record["pqu_id"],
                    "change_type": "baseline",
                    "previous_status": None,
                    "new_status": record["status"],
                    "source_commit": record["source"]["commit"],
                }
            )
    return rows


def generate_workbook(
    path: Path,
    *,
    dataset: NormalizedDataset,
    metadata: dict[str, Any],
    health: dict[str, Any],
    quality: Mapping[str, Any],
) -> None:
    wb = Workbook()
    active = wb.active
    if active is not None:
        wb.remove(active)
    _readme(wb, dataset, metadata, health)

    pqu_columns: list[Column] = [
        ("PQU ID", "pqu_id", 18),
        ("Application Version", "application_version", 18),
        ("PQU Train", "pqu_train", 12),
        ("Release Number", "release_number", 14),
        ("Change Cutoff Date", "change_cutoff_date", 20),
        ("Train Start Date", "train_start_date", 16),
        ("Train End Date", "train_end_date", 16),
        ("Status", "status", 14),
        ("Application Build", "application_build", 20),
        ("Platform Build", "platform_build", 18),
        ("UEP Version", "uep_version", 14),
        ("Station Schedule", "station_schedule_available", 16),
        ("Source", "source.publisher", 12),
        ("Source URL", "source.url", 46),
        ("Source Commit", "source.commit", 42),
        ("Last Updated", "last_changed_at", 24),
        ("First Seen", "first_seen_at", 24),
        ("Last Changed", "last_changed_at", 24),
    ]
    _write_table_sheet(wb, "01_PQU_MASTER", pqu_columns, dataset.records, table_name="PquMaster")
    _write_table_sheet(
        wb,
        "02_CURRENT_PQU",
        pqu_columns,
        [r for r in dataset.records if r["status"] == "In-Progress"],
        table_name="CurrentPqu",
    )
    _write_table_sheet(
        wb,
        "03_UPCOMING_PQU",
        pqu_columns,
        [r for r in dataset.records if r["status"] == "Not Started"],
        table_name="UpcomingPqu",
    )

    station_columns: list[Column] = [
        ("PQU ID", "pqu_id", 18),
        ("Application Version", "application_version", 18),
        ("PQU Train", "pqu_train", 12),
        ("Station", "station", 10),
        ("Station Label", "station_label", 14),
        ("Sandbox Start", "sandbox_start_date", 16),
        ("Sandbox End", "sandbox_end_date", 16),
        ("Production Start", "production_start_date", 18),
        ("Production End", "production_end_date", 18),
        ("Source Commit", "source.commit", 42),
    ]
    _write_table_sheet(
        wb, "04_STATION_SCHEDULE", station_columns, dataset.stations, table_name="StationSchedule"
    )
    region_columns: list[Column] = [
        ("Station", "station", 10),
        ("Station Label", "station_label", 14),
        ("Region", "region", 28),
        ("Is Region", "is_region", 12),
    ]
    _write_table_sheet(
        wb, "05_REGION_MAPPING", region_columns, dataset.regions, table_name="RegionMapping"
    )
    _write_table_sheet(
        wb,
        "06_APPLICATION_BUILDS",
        [
            ("PQU ID", "pqu_id", 18),
            ("Application Version", "application_version", 18),
            ("Application Build", "application_build", 20),
            ("Status", "status", 14),
            ("Change Cutoff Date", "change_cutoff_date", 20),
            ("Train Start Date", "train_start_date", 16),
            ("Train End Date", "train_end_date", 16),
        ],
        dataset.versions,
        table_name="ApplicationBuilds",
    )
    _write_table_sheet(
        wb,
        "07_PLATFORM_BUILDS",
        [
            ("PQU ID", "pqu_id", 18),
            ("Application Version", "application_version", 18),
            ("Platform Build", "platform_build", 18),
            ("Status", "status", 14),
            ("Change Cutoff Date", "change_cutoff_date", 20),
            ("Train Start Date", "train_start_date", 16),
            ("Train End Date", "train_end_date", 16),
        ],
        dataset.versions,
        table_name="PlatformBuilds",
    )
    _write_table_sheet(
        wb,
        "08_UEP_BUILDS",
        [
            ("PQU ID", "pqu_id", 18),
            ("Application Version", "application_version", 18),
            ("UEP Version", "uep_version", 14),
            ("Status", "status", 14),
            ("Change Cutoff Date", "change_cutoff_date", 20),
            ("Train Start Date", "train_start_date", 16),
            ("Train End Date", "train_end_date", 16),
        ],
        dataset.versions,
        table_name="UepBuilds",
    )
    _write_table_sheet(
        wb,
        "09_STATUS_HISTORY",
        [
            ("Changed At", "changed_at", 24),
            ("PQU ID", "pqu_id", 18),
            ("Change Type", "change_type", 14),
            ("Previous Status", "previous_status", 16),
            ("New Status", "new_status", 16),
            ("Source Commit", "source_commit", 42),
        ],
        _status_history_rows(dataset, dataset.generated_at),
        table_name="StatusHistory",
    )
    _write_table_sheet(
        wb,
        "10_CHANGE_HISTORY",
        [
            ("Change ID", "change_id", 24),
            ("Changed At", "changed_at", 24),
            ("PQU ID", "pqu_id", 18),
            ("Entity", "entity", 12),
            ("Change Type", "change_type", 14),
            ("Field", "field", 22),
            ("Old Value", "old_value", 34),
            ("New Value", "new_value", 34),
            ("Source Commit", "source_commit", 42),
        ],
        dataset.changes,
        table_name="ChangeHistory",
    )
    _key_value_sheet(
        wb,
        "11_SOURCE_METADATA",
        [
            ("Dataset", DATASET_NAME),
            ("Schema Version", SCHEMA_VERSION),
            ("Pipeline Revision", metadata.get("pipeline_revision", "")),
            ("Publisher", dataset.source["publisher"]),
            ("Source Repository", dataset.source["repository"]),
            ("Source Branch", dataset.source["branch"]),
            ("Source File", dataset.source["file_path"]),
            ("Source Commit", dataset.source["commit"]),
            ("Source Article", dataset.source["article_url"]),
            ("Raw Source", dataset.source["raw_url"]),
            ("Source Markdown Date", dataset.source["markdown_date"]),
            ("Source SHA-256", dataset.source["sha256"]),
            ("Retrieved At", dataset.source["retrieved_at"]),
            ("Generated At", dataset.generated_at.isoformat()),
            ("Hours Between Checks", 6),
            ("Workbook", metadata["links"]["workbook"]),
        ],
    )
    _key_value_sheet(
        wb,
        "12_SYNC_HEALTH",
        [
            ("Status", health["status"]),
            ("Checked At", health["checked_at"]),
            ("Last Successful Check", health["last_successful_check_at"]),
            ("Last Successful Publish", health["last_successful_publish_at"]),
            ("Last Source Change", health["last_source_change_at"]),
            ("Source Reachable", health["source_reachable"]),
            ("Parser Success", health["parser_success"]),
            ("Validation Passed", health["validation_passed"]),
            ("Records", health["records"]),
            ("Active Records", health["current_records"]),
            ("Upcoming Records", health["upcoming_records"]),
            ("Warnings", health["warning_count"]),
            ("Errors", health["error_count"]),
            ("Source Commit", health["source_commit"]),
            ("Source SHA-256", health["source_sha256"]),
        ],
    )
    _write_table_sheet(
        wb,
        "13_DATA_QUALITY",
        [
            ("Code", "code", 28),
            ("Severity", "severity", 12),
            ("Message", "message", 90),
            ("PQU ID", "pqu_id", 18),
            ("Field", "field", 24),
        ],
        quality.get("records", []),
        table_name="DataQuality",
    )

    wb.properties.title = "D365 Finance & Operations PQU Tracker"
    wb.properties.subject = "Microsoft Dynamics 365 proactive quality update schedule"
    wb.properties.creator = "d365-pqu-data automation"
    wb.properties.lastModifiedBy = "d365-pqu-data automation"
    wb.properties.description = (
        f"Generated {metadata['generated_at']} from Microsoft public documentation."
    )
    wb.properties.keywords = (
        "D365, Finance and Operations, PQU, proactive quality update, Microsoft"
    )
    created = datetime.fromisoformat(metadata["generated_at"].replace("Z", "+00:00"))
    wb.properties.created = created
    wb.properties.modified = created
    wb.active = 0
    ws = wb["00_README"]
    ws.sheet_properties.tabColor = "0B3D5C"
    for name in ("01_PQU_MASTER", "02_CURRENT_PQU", "03_UPCOMING_PQU"):
        wb[name].sheet_properties.tabColor = "12506E"
    for name in ("04_STATION_SCHEDULE", "05_REGION_MAPPING"):
        wb[name].sheet_properties.tabColor = "2E7D5B"
    for name in ("06_APPLICATION_BUILDS", "07_PLATFORM_BUILDS", "08_UEP_BUILDS"):
        wb[name].sheet_properties.tabColor = "7A5AA6"
    for name in ("09_STATUS_HISTORY", "10_CHANGE_HISTORY"):
        wb[name].sheet_properties.tabColor = "B5651D"
    for name in ("11_SOURCE_METADATA", "12_SYNC_HEALTH", "13_DATA_QUALITY"):
        wb[name].sheet_properties.tabColor = "6B7A8F"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def now_utc() -> datetime:
    return datetime.now(UTC)
