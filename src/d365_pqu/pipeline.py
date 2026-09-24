from __future__ import annotations

import json
import os
import shutil
import uuid
from collections.abc import Mapping
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from d365_pqu.config import HEARTBEAT_DAYS, PIPELINE_REVISION, Paths
from d365_pqu.diff import compute_changes
from d365_pqu.errors import PublishError, ValidationError
from d365_pqu.excel import generate_workbook
from d365_pqu.models import NormalizedDataset, SourceDocument, SyncResult
from d365_pqu.normalize import normalize_source
from d365_pqu.parser import parse_source
from d365_pqu.serialization import (
    CHANGE_CSV_FIELDS,
    PQU_CSV_FIELDS,
    QUALITY_CSV_FIELDS,
    REGION_CSV_FIELDS,
    STATION_CSV_FIELDS,
    VERSION_CSV_FIELDS,
    changes_document,
    current_document,
    envelope,
    flatten_pqu_record,
    flatten_station_record,
    health_document,
    metadata_document,
    quality_document,
    upcoming_document,
    write_csv,
    write_json,
)
from d365_pqu.source import fetch_source, parse_markdown_date
from d365_pqu.validation import (
    validate_document,
    validate_quality,
    validate_records,
)

EXPECTED_SHEETS = (
    "00_README",
    "01_PQU_MASTER",
    "02_CURRENT_PQU",
    "03_UPCOMING_PQU",
    "04_STATION_SCHEDULE",
    "05_REGION_MAPPING",
    "06_APPLICATION_BUILDS",
    "07_PLATFORM_BUILDS",
    "08_UEP_BUILDS",
    "09_STATUS_HISTORY",
    "10_CHANGE_HISTORY",
    "11_SOURCE_METADATA",
    "12_SYNC_HEALTH",
    "13_DATA_QUALITY",
)

SCHEMA_FOR_DOCUMENT = {
    "pqu.json": "pqu.schema.json",
    "pqu-current.json": "pqu.schema.json",
    "pqu-upcoming.json": "pqu.schema.json",
    "pqu-stations.json": "station.schema.json",
    "pqu-regions.json": "region.schema.json",
    "pqu-versions.json": "version.schema.json",
    "pqu-changes.json": "change.schema.json",
    "quality-report.json": "quality-report.schema.json",
    "metadata.json": "metadata.schema.json",
    "health.json": "health.schema.json",
}


def load_json(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValidationError(f"{path} is not valid JSON: {exc}") from exc


def load_previous(
    paths: Paths,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None, dict[str, Any] | None]:
    pqu = load_json(paths.pqu_path)
    metadata = load_json(paths.metadata_path)
    health = load_json(paths.health_path)
    return pqu, metadata, health


def previous_records(pqu: dict[str, Any] | None) -> list[dict[str, Any]] | None:
    if not pqu:
        return None
    records = pqu.get("records")
    if not isinstance(records, list):
        return None
    return [record for record in records if isinstance(record, dict)]


def source_document_from_file(
    path: Path,
    *,
    commit: str,
    now: datetime,
) -> SourceDocument:
    import hashlib

    markdown = path.read_text(encoding="utf-8")
    sha256 = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    return SourceDocument(
        markdown=markdown,
        source_commit=commit,
        article_url="https://learn.microsoft.com/en-us/dynamics365/fin-ops-core/dev-itpro/get-started/quality-updates-schedule",
        raw_url=f"https://raw.githubusercontent.com/MicrosoftDocs/dynamics-365-unified-operations-public/{commit}/articles/fin-ops-core/dev-itpro/get-started/quality-updates-schedule.md",
        markdown_date=parse_markdown_date(markdown),
        retrieved_at=now,
        sha256=sha256,
        warnings=(),
    )


def heartbeat_due(health: dict[str, Any] | None, now: datetime) -> bool:
    if not health:
        return True
    checked = health.get("checked_at")
    if not isinstance(checked, str):
        return True
    try:
        previous = datetime.fromisoformat(checked.replace("Z", "+00:00"))
    except ValueError:
        return True
    if previous.tzinfo is None:
        previous = previous.replace(tzinfo=UTC)
    return now - previous >= timedelta(days=HEARTBEAT_DAYS)


def _atomic_write_json(path: Path, document: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    write_json(temporary, document)
    os.replace(temporary, path)


def _write_artifacts(
    paths: Paths,
    *,
    dataset: NormalizedDataset,
    documents: dict[str, dict[str, Any]],
    metadata: dict[str, Any],
    health: dict[str, Any],
    quality: Mapping[str, Any],
) -> None:
    staging = paths.root / f".d365-pqu-staging-{uuid.uuid4().hex}"
    staged_paths = Paths(
        root=staging,
        data_dir=staging / "data",
        excel_dir=staging / "excel",
        schema_dir=paths.schema_dir,
        site_dir=staging / "site",
        tests_dir=paths.tests_dir,
    )
    try:
        if (paths.data_dir / "history").exists():
            shutil.copytree(
                paths.data_dir / "history",
                staged_paths.data_dir / "history",
                dirs_exist_ok=True,
            )
        _write_artifacts_in_place(
            staged_paths,
            dataset=dataset,
            documents=documents,
            metadata=metadata,
            health=health,
            quality=quality,
        )
        verify_artifacts(staged_paths)
        _swap_directories(
            (staged_paths.data_dir, paths.data_dir),
            (staged_paths.excel_dir, paths.excel_dir),
        )
    finally:
        shutil.rmtree(staging, ignore_errors=True)


def _swap_directories(*pairs: tuple[Path, Path]) -> None:
    backups: list[tuple[Path, Path | None]] = []
    try:
        for source, target in pairs:
            backup = target.with_name(f".{target.name}.backup-{uuid.uuid4().hex}")
            if target.exists():
                os.replace(target, backup)
                backups.append((target, backup))
            else:
                backups.append((target, None))
            os.replace(source, target)
        for _, saved in backups:
            if saved is not None and saved.exists():
                shutil.rmtree(saved)
    except Exception:
        for target, saved in reversed(backups):
            if target.exists():
                shutil.rmtree(target)
            if saved is not None and saved.exists():
                os.replace(saved, target)
        raise


def _write_artifacts_in_place(
    paths: Paths,
    *,
    dataset: NormalizedDataset,
    documents: dict[str, dict[str, Any]],
    metadata: dict[str, Any],
    health: dict[str, Any],
    quality: Mapping[str, Any],
) -> None:
    for name, document in documents.items():
        _atomic_write_json(paths.data_dir / name, document)

    records = dataset.records
    write_csv(
        paths.data_dir / "pqu.csv",
        (flatten_pqu_record(record) for record in records),
        PQU_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-current.csv",
        (flatten_pqu_record(r) for r in records if r["status"] == "In-Progress"),
        PQU_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-upcoming.csv",
        (flatten_pqu_record(r) for r in records if r["status"] == "Not Started"),
        PQU_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-stations.csv",
        (flatten_station_record(record) for record in dataset.stations),
        STATION_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-regions.csv",
        (dict(record) for record in dataset.regions),
        REGION_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-versions.csv",
        (dict(record) for record in dataset.versions),
        VERSION_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "pqu-changes.csv",
        (dict(record) for record in dataset.changes),
        CHANGE_CSV_FIELDS,
    )
    write_csv(
        paths.data_dir / "quality-report.csv",
        (dict(record) for record in quality.get("records", [])),
        QUALITY_CSV_FIELDS,
    )

    history_dir = (
        paths.data_dir
        / "history"
        / f"{dataset.generated_at.year:04d}"
        / f"{dataset.generated_at.month:02d}"
    )
    history_name = (
        f"{dataset.generated_at.strftime('%Y-%m-%dT%H%M%S%fZ')}-"
        f"{dataset.source['commit'][:12]}.json"
    )
    _atomic_write_json(history_dir / history_name, documents["pqu.json"])

    generate_workbook(
        paths.workbook_path,
        dataset=dataset,
        metadata=metadata,
        health=health,
        quality=quality,
    )


def build_documents(
    dataset: NormalizedDataset,
    previous_metadata: dict[str, Any] | None,
    previous_health: dict[str, Any] | None,
    changed: bool,
) -> dict[str, dict[str, Any]]:
    source = dict(dataset.source)
    metadata = metadata_document(
        dataset,
        previous_metadata=previous_metadata,
        generated_at=dataset.generated_at,
    )
    health = health_document(
        dataset,
        previous_health=previous_health,
        checked_at=dataset.generated_at,
        changed=changed,
    )
    quality = quality_document(
        dataset.quality,
        source=source,
        generated_at=dataset.generated_at,
    )
    documents = {
        "pqu.json": envelope(
            [dict(record) for record in dataset.records],
            source=source,
            generated_at=dataset.generated_at,
        ),
        "pqu-current.json": current_document(dataset),
        "pqu-upcoming.json": upcoming_document(dataset),
        "pqu-stations.json": envelope(
            [dict(record) for record in dataset.stations],
            source=source,
            generated_at=dataset.generated_at,
        ),
        "pqu-regions.json": envelope(
            [dict(record) for record in dataset.regions],
            source=source,
            generated_at=dataset.generated_at,
        ),
        "pqu-versions.json": envelope(
            [dict(record) for record in dataset.versions],
            source=source,
            generated_at=dataset.generated_at,
        ),
        "pqu-changes.json": changes_document(
            [dict(record) for record in dataset.changes],
            source=source,
            generated_at=dataset.generated_at,
        ),
        "quality-report.json": quality,
        "metadata.json": metadata,
        "health.json": health,
    }
    return documents


def validate_documents(paths: Paths, documents: dict[str, dict[str, Any]]) -> None:
    for name, document in documents.items():
        schema_name = SCHEMA_FOR_DOCUMENT[name]
        validate_document(document, paths.schema_dir / schema_name)


def _unchanged_result(
    paths: Paths,
    *,
    source_hash: str,
    source_commit: str,
    previous_metadata: dict[str, Any] | None,
    heartbeat: bool,
) -> SyncResult:
    metadata = previous_metadata or {}
    quality = load_json(paths.data_dir / "quality-report.json") or {}
    warnings = int(quality.get("warning_count", 0)) if isinstance(quality, dict) else 0
    return SyncResult(
        changed=heartbeat,
        status="heartbeat" if heartbeat else "unchanged",
        source_hash=source_hash,
        source_commit=source_commit,
        record_count=int(metadata.get("record_count", 0)),
        current_count=int(metadata.get("current_count", 0)),
        upcoming_count=int(metadata.get("upcoming_count", 0)),
        warning_count=warnings,
        error_count=0,
        data_dir=str(paths.data_dir),
        workbook_path=str(paths.workbook_path),
        site_dir=str(paths.site_dir),
    )


def run_sync(
    paths: Paths,
    *,
    source: SourceDocument | None = None,
    now: datetime | None = None,
    write: bool = True,
) -> SyncResult:
    moment = now or datetime.now(UTC)
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=UTC)
    pqu, previous_metadata, previous_health = load_previous(paths)
    document = source or fetch_source(now=moment)
    parsed = parse_source(document.markdown)
    dataset = normalize_source(parsed, document, previous=pqu, observed_at=moment)
    validate_quality(dataset)
    validate_records(dataset)

    previous_source = (previous_metadata or {}).get("source") if previous_metadata else None
    previous_hash = previous_source.get("sha256") if isinstance(previous_source, dict) else None
    previous_commit = previous_source.get("commit") if isinstance(previous_source, dict) else None
    identity_unchanged = bool(
        pqu
        and previous_hash == document.sha256
        and previous_commit == document.source_commit
        and (previous_metadata or {}).get("pipeline_revision") == PIPELINE_REVISION
    )
    heartbeat = heartbeat_due(previous_health, moment)
    if identity_unchanged and not heartbeat:
        return _unchanged_result(
            paths,
            source_hash=document.sha256,
            source_commit=document.source_commit,
            previous_metadata=previous_metadata,
            heartbeat=False,
        )
    if identity_unchanged and heartbeat and write:
        _refresh_health_only(
            paths,
            dataset=dataset,
            previous_metadata=previous_metadata,
            previous_health=previous_health,
            now=moment,
        )
        return _unchanged_result(
            paths,
            source_hash=document.sha256,
            source_commit=document.source_commit,
            previous_metadata=previous_metadata,
            heartbeat=True,
        )

    previous_change_document = load_json(paths.data_dir / "pqu-changes.json")
    prior_changes = (
        previous_change_document.get("records", [])
        if previous_change_document and isinstance(previous_change_document.get("records"), list)
        else []
    )
    previous_station_document = load_json(paths.data_dir / "pqu-stations.json")
    previous_region_document = load_json(paths.data_dir / "pqu-regions.json")
    previous_stations = (
        previous_station_document.get("records", [])
        if previous_station_document and isinstance(previous_station_document.get("records"), list)
        else []
    )
    previous_regions = (
        previous_region_document.get("records", [])
        if previous_region_document and isinstance(previous_region_document.get("records"), list)
        else []
    )
    new_changes = compute_changes(
        previous_records(pqu),
        [dict(record) for record in dataset.records],
        source_commit=document.source_commit,
        changed_at=moment,
        previous_stations=previous_stations,
        current_stations=[dict(record) for record in dataset.stations],
        previous_regions=previous_regions,
        current_regions=[dict(record) for record in dataset.regions],
    )
    changed_station_ids = {
        change["pqu_id"] for change in new_changes if change["entity"] == "station"
    }
    if changed_station_ids:
        changed_at = moment.isoformat().replace("+00:00", "Z")
        for record in dataset.records:
            if record["pqu_id"] in changed_station_ids:
                record["last_changed_at"] = changed_at
    dataset.changes = [*prior_changes, *new_changes]
    validate_quality(dataset)
    validate_records(dataset)

    documents = build_documents(dataset, previous_metadata, previous_health, True)
    validate_documents(paths, documents)

    if write:
        _write_artifacts(
            paths,
            dataset=dataset,
            documents=documents,
            metadata=documents["metadata.json"],
            health=documents["health.json"],
            quality=documents["quality-report.json"],
        )

    return SyncResult(
        changed=True,
        status="updated",
        source_hash=document.sha256,
        source_commit=document.source_commit,
        record_count=len(dataset.records),
        current_count=sum(1 for record in dataset.records if record["status"] == "In-Progress"),
        upcoming_count=sum(1 for record in dataset.records if record["status"] == "Not Started"),
        warning_count=sum(1 for item in dataset.quality if item["severity"] == "warning"),
        error_count=0,
        data_dir=str(paths.data_dir),
        workbook_path=str(paths.workbook_path),
        site_dir=str(paths.site_dir),
    )


def _refresh_health_only(
    paths: Paths,
    *,
    dataset: NormalizedDataset,
    previous_metadata: dict[str, Any] | None,
    previous_health: dict[str, Any] | None,
    now: datetime,
) -> None:
    validate_quality(dataset)
    validate_records(dataset)
    health = health_document(
        dataset,
        previous_health=previous_health,
        checked_at=now,
        changed=False,
    )
    validate_document(health, paths.schema_dir / "health.schema.json")
    _atomic_write_json(paths.health_path, health)


def build_site(paths: Paths) -> None:
    from d365_pqu.site import generate_site

    metadata = load_json(paths.metadata_path)
    if not metadata:
        raise PublishError("metadata.json is missing; run sync before building the site")
    generate_site(
        site_dir=paths.site_dir,
        data_dir=paths.data_dir,
        schema_dir=paths.schema_dir,
        workbook_path=paths.workbook_path,
        metadata=metadata,
        static_dir=Path(__file__).parent / "static",
    )


def _document_records(document: dict[str, Any], name: str) -> list[dict[str, Any]]:
    records = document.get("records")
    if not isinstance(records, list) or not all(isinstance(record, dict) for record in records):
        raise ValidationError(f"{name} has an invalid records collection")
    return records


def _verify_csv(
    path: Path,
    *,
    fieldnames: tuple[str, ...],
    expected_count: int,
) -> None:
    import csv

    if not path.exists():
        raise ValidationError(f"Missing generated CSV {path}")
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if tuple(reader.fieldnames or ()) != fieldnames:
            raise ValidationError(
                f"{path.name} headers do not match the generated schema: {reader.fieldnames}"
            )
        rows = list(reader)
    if len(rows) != expected_count:
        raise ValidationError(f"{path.name} has {len(rows)} rows; expected {expected_count}")


def verify_artifacts(paths: Paths) -> dict[str, Any]:
    documents: dict[str, dict[str, Any]] = {}
    for name, schema_name in SCHEMA_FOR_DOCUMENT.items():
        document = load_json(paths.data_dir / name)
        if document is None:
            raise ValidationError(f"Missing required artifact {paths.data_dir / name}")
        validate_document(document, paths.schema_dir / schema_name)
        documents[name] = document

    pqu_records = _document_records(documents["pqu.json"], "pqu.json")
    pqu_ids = [record["pqu_id"] for record in pqu_records]
    if len(pqu_ids) != len(set(pqu_ids)):
        raise ValidationError("pqu.json contains duplicate identifiers")
    pqu_by_id = {record["pqu_id"]: record for record in pqu_records}
    current_records = _document_records(documents["pqu-current.json"], "pqu-current.json")
    upcoming_records = _document_records(documents["pqu-upcoming.json"], "pqu-upcoming.json")
    expected_current = [record for record in pqu_records if record["status"] == "In-Progress"]
    expected_upcoming = [record for record in pqu_records if record["status"] == "Not Started"]
    if [record["pqu_id"] for record in current_records] != [
        record["pqu_id"] for record in expected_current
    ]:
        raise ValidationError("pqu-current.json does not contain the In-Progress records")
    if [record["pqu_id"] for record in upcoming_records] != [
        record["pqu_id"] for record in expected_upcoming
    ]:
        raise ValidationError("pqu-upcoming.json does not contain the Not Started records")
    if documents["pqu-current.json"].get("latest_pqu_id") != documents["metadata.json"].get(
        "latest_pqu_id"
    ):
        raise ValidationError("Current latest_pqu_id does not match metadata.json")

    source = documents["pqu.json"]["source"]
    for name, document in documents.items():
        if name == "health.json":
            continue
        if document["source"] != source:
            raise ValidationError(f"Source provenance differs in {name}")
    metadata = documents["metadata.json"]
    if metadata.get("pipeline_revision") != PIPELINE_REVISION:
        raise ValidationError("metadata.json pipeline revision is stale")
    if metadata.get("record_count") != len(pqu_records):
        raise ValidationError("metadata record_count does not match pqu.json")
    if metadata.get("current_count") != len(current_records):
        raise ValidationError("metadata current_count does not match pqu-current.json")
    if metadata.get("upcoming_count") != len(upcoming_records):
        raise ValidationError("metadata upcoming_count does not match pqu-upcoming.json")

    station_records = _document_records(documents["pqu-stations.json"], "pqu-stations.json")
    region_records = _document_records(documents["pqu-regions.json"], "pqu-regions.json")
    version_records = _document_records(documents["pqu-versions.json"], "pqu-versions.json")
    change_records = _document_records(documents["pqu-changes.json"], "pqu-changes.json")
    quality_records = _document_records(documents["quality-report.json"], "quality-report.json")
    if len(station_records) != metadata.get("station_schedule_count"):
        raise ValidationError("metadata station_schedule_count does not match stations")
    if len(region_records) != metadata.get("region_count"):
        raise ValidationError("metadata region_count does not match regions")
    if len(quality_records) != documents["quality-report.json"].get("count"):
        raise ValidationError("quality-report count does not match records")
    if any(record["pqu_id"] not in pqu_by_id for record in station_records):
        raise ValidationError("Station schedule references an unknown PQU")
    if any(record["pqu_id"] not in pqu_by_id for record in version_records):
        raise ValidationError("Version record references an unknown PQU")
    if any(change.get("entity") not in {"pqu", "station", "region"} for change in change_records):
        raise ValidationError("Change history contains an unknown entity")
    for record in station_records:
        if not pqu_by_id[record["pqu_id"]].get("station_schedule_available"):
            raise ValidationError(
                f"Station schedule is not marked available for {record['pqu_id']}"
            )

    expected_csvs = {
        "pqu.csv": (PQU_CSV_FIELDS, len(pqu_records)),
        "pqu-current.csv": (PQU_CSV_FIELDS, len(current_records)),
        "pqu-upcoming.csv": (PQU_CSV_FIELDS, len(upcoming_records)),
        "pqu-stations.csv": (STATION_CSV_FIELDS, len(station_records)),
        "pqu-regions.csv": (REGION_CSV_FIELDS, len(region_records)),
        "pqu-versions.csv": (VERSION_CSV_FIELDS, len(version_records)),
        "pqu-changes.csv": (CHANGE_CSV_FIELDS, len(change_records)),
        "quality-report.csv": (QUALITY_CSV_FIELDS, len(quality_records)),
    }
    for name, (fields, count) in expected_csvs.items():
        _verify_csv(paths.data_dir / name, fieldnames=fields, expected_count=count)

    workbook = paths.workbook_path
    if not workbook.exists():
        raise ValidationError(f"Missing workbook {workbook}")
    from openpyxl import load_workbook

    wb = load_workbook(workbook, read_only=True)
    missing = [sheet for sheet in EXPECTED_SHEETS if sheet not in wb.sheetnames]
    if missing:
        raise ValidationError(f"Workbook is missing sheets: {missing}")
    expected_sheet_rows = {
        "01_PQU_MASTER": len(pqu_records),
        "02_CURRENT_PQU": len(current_records),
        "03_UPCOMING_PQU": len(upcoming_records),
        "04_STATION_SCHEDULE": len(station_records),
        "05_REGION_MAPPING": len(region_records),
        "06_APPLICATION_BUILDS": len(version_records),
        "07_PLATFORM_BUILDS": len(version_records),
        "08_UEP_BUILDS": len(version_records),
        "10_CHANGE_HISTORY": len(change_records),
        "13_DATA_QUALITY": len(quality_records),
    }
    for sheet, expected in expected_sheet_rows.items():
        actual = wb[sheet].max_row - 1
        if actual != expected:
            raise ValidationError(
                f"Workbook sheet {sheet} has {actual} data rows; expected {expected}"
            )
    wb.close()
    return {
        "status": "verified",
        "record_count": len(pqu_records),
        "latest_pqu_id": metadata.get("latest_pqu_id"),
        "semantic_hash": semantic_hash_from_document({"records": pqu_records}),
    }


def semantic_hash_from_document(document: dict[str, Any]) -> str:
    import hashlib

    payload = json.dumps(
        document.get("records", []), sort_keys=True, separators=(",", ":"), default=str
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def remove_site(paths: Paths) -> None:
    if paths.site_dir.exists():
        shutil.rmtree(paths.site_dir)
