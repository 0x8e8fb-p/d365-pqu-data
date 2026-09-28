from __future__ import annotations

import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import LIVE_FIXTURES, SCHEMA_DIR
from d365_pqu.config import Paths
from d365_pqu.errors import ValidationError
from d365_pqu.models import SourceBundle
from d365_pqu.pipeline import load_json, run_sync, verify_artifacts
from d365_pqu.source import bundle_from_directory

SERVICE_FILE = "public-preview-releases.md"
DAY_1 = datetime(2026, 9, 24, 12, 0, tzinfo=UTC)
DAY_2 = datetime(2026, 9, 25, 12, 0, tzinfo=UTC)
DAY_3 = datetime(2026, 9, 26, 12, 0, tzinfo=UTC)


def _paths(tmp_path: Path) -> Paths:
    return Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        excel_dir=tmp_path / "excel",
        schema_dir=SCHEMA_DIR,
        site_dir=tmp_path / "build" / "site",
        tests_dir=tmp_path / "tests",
    )


def _sources(tmp_path: Path, name: str = "sources") -> Path:
    directory = tmp_path / name
    shutil.copytree(LIVE_FIXTURES, directory)
    return directory


def _sync(paths: Paths, directory: Path, commit: str, now: datetime, **errors: str):
    bundle = bundle_from_directory(directory, commit=commit, now=now)
    if errors:
        documents = {key: value for key, value in bundle.documents.items() if key not in errors}
        bundle = SourceBundle(commit=commit, documents=documents, errors=dict(errors))
    return run_sync(paths, bundle=bundle, now=now)


def _quality_codes(paths: Paths) -> list[str]:
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    return [item["code"] for item in quality["records"]]


def test_all_sources_publish_service_updates_with_their_own_provenance(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    result = _sync(paths, _sources(tmp_path), "a" * 40, DAY_1)
    assert result.source_states == {
        "schedule": "updated",
        "service_updates": "updated",
        "maintenance": "updated",
        "pqu_overview": "updated",
        "pqu_faq": "updated",
    }
    document = load_json(paths.data_dir / "service-updates.json")
    assert document is not None
    assert document["state"] == "current"
    assert document["count"] == 7
    assert document["source"]["file_path"].endswith("get-started/public-preview-releases.md")
    assert document["source"]["markdown_date"] == "2026-06-04"
    assert document["source"]["commit"] == "a" * 40
    metadata = load_json(paths.metadata_path)
    assert metadata is not None
    assert metadata["service_update_count"] == 7
    assert metadata["sources"]["schedule"]["source"] == metadata["source"]
    assert metadata["sources"]["service_updates"]["source"] == document["source"]
    assert metadata["sources"]["service_updates"]["label"] == "Service update availability"
    learn = load_json(paths.data_dir / "learn.json")
    assert learn is not None
    assert [article["key"] for article in learn["articles"]] == [
        "schedule",
        "service_updates",
        "maintenance",
        "pqu_overview",
        "pqu_faq",
    ]
    service_article = learn["articles"][1]
    assert service_article["markdown_date"] == "2026-06-04"
    placeholders = [
        block["name"]
        for section in service_article["sections"]
        for block in section["blocks"]
        if block["type"] == "dataset"
    ]
    assert placeholders == ["service_updates"]
    assert "source-stale" not in _quality_codes(paths)
    assert verify_artifacts(paths)["status"] == "verified"
    workbook = load_workbook(paths.workbook_path, read_only=True)
    assert workbook["14_SERVICE_UPDATES"].max_row == 8
    workbook.close()
    csv_lines = (paths.data_dir / "service-updates.csv").read_text(encoding="utf-8").splitlines()
    assert len(csv_lines) == 8
    assert csv_lines[0].startswith("version,release_label,is_major,")


def test_schedule_only_runs_leave_optional_sources_not_configured(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = tmp_path / "schedule-only"
    directory.mkdir()
    shutil.copy(LIVE_FIXTURES / "quality-updates-schedule.md", directory)
    result = _sync(paths, directory, "a" * 40, DAY_1)
    assert result.source_states["service_updates"] == "not_configured"
    assert result.source_states["maintenance"] == "not_configured"
    document = load_json(paths.data_dir / "service-updates.json")
    assert document is not None
    assert (document["state"], document["source"], document["count"]) == ("not_configured", None, 0)
    regions = load_json(paths.data_dir / "pqu-regions.json")
    assert regions is not None
    assert {record["maintenance_geo"] for record in regions["records"]} == {None}
    assert _quality_codes(paths) == ["cutoff-start-year-mismatch"]
    assert verify_artifacts(paths)["status"] == "verified"


def test_broken_optional_article_keeps_last_good_data_while_schedule_updates(
    tmp_path: Path,
) -> None:
    paths = _paths(tmp_path)
    directory = _sources(tmp_path)
    _sync(paths, directory, "a" * 40, DAY_1)
    before = load_json(paths.data_dir / "service-updates.json")
    service = directory / SERVICE_FILE
    service.write_text(
        service.read_text(encoding="utf-8").replace("| Release version |", "| Release |"),
        encoding="utf-8",
    )
    schedule = directory / "quality-updates-schedule.md"
    schedule.write_text(
        schedule.read_text(encoding="utf-8").replace("7.0.7996.119", "7.0.7996.130"),
        encoding="utf-8",
    )
    result = _sync(paths, directory, "b" * 40, DAY_2)
    assert result.status == "updated"
    assert result.source_states == {
        "schedule": "updated",
        "service_updates": "stale",
        "maintenance": "unchanged",
        "pqu_overview": "unchanged",
        "pqu_faq": "unchanged",
    }
    after = load_json(paths.data_dir / "service-updates.json")
    assert before is not None and after is not None
    assert after["state"] == "stale"
    assert after["records"] == before["records"]
    assert after["source"] == before["source"]
    assert after["source"]["commit"] == "a" * 40
    pqu = load_json(paths.pqu_path)
    assert pqu is not None
    assert pqu["source"]["commit"] == "b" * 40
    record = next(r for r in pqu["records"] if r["pqu_id"] == "10.0.48-PQU-6")
    assert record["platform_build"] == "7.0.7996.130"
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    stale = [item for item in quality["records"] if item["code"] == "source-stale"]
    assert len(stale) == 1
    assert stale[0]["severity"] == "warning"
    assert stale[0]["message"].startswith("Service update availability could not be refreshed")
    assert f"commit {'a' * 12}" in stale[0]["message"]
    learn = load_json(paths.data_dir / "learn.json")
    assert learn is not None
    kept = next(article for article in learn["articles"] if article["key"] == "service_updates")
    assert kept["commit"] == "a" * 40
    run_health = load_json(paths.run_health_path)
    assert run_health is not None
    assert run_health["sources"]["service_updates"]["state"] == "stale"
    assert "headers changed" in run_health["sources"]["service_updates"]["message"]
    assert verify_artifacts(paths)["status"] == "verified"

    # A later run with the same broken article does not rewrite anything.
    unchanged = _sync(paths, directory, "c" * 40, DAY_3)
    assert unchanged.status == "unchanged"

    # Recovery publishes the fresh article again and clears the warning.
    shutil.copy(LIVE_FIXTURES / SERVICE_FILE, service)
    recovered = _sync(paths, directory, "d" * 40, DAY_3)
    assert recovered.status == "updated"
    assert recovered.source_states["service_updates"] == "updated"
    current = load_json(paths.data_dir / "service-updates.json")
    assert current is not None
    assert current["state"] == "current"
    assert current["source"]["commit"] == "d" * 40
    assert "source-stale" not in _quality_codes(paths)


def test_unreachable_optional_article_without_history_is_unavailable(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    result = _sync(
        paths, _sources(tmp_path), "a" * 40, DAY_1, service_updates="HTTP Error 503: Unavailable"
    )
    assert result.source_states["service_updates"] == "unavailable"
    document = load_json(paths.data_dir / "service-updates.json")
    assert document is not None
    assert (document["state"], document["source"], document["records"]) == (
        "unavailable",
        None,
        [],
    )
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    unavailable = [item for item in quality["records"] if item["code"] == "source-unavailable"]
    assert unavailable[0]["message"] == (
        "Service update availability could not be retrieved (HTTP Error 503: Unavailable); "
        "information from it is not shown."
    )
    assert verify_artifacts(paths)["status"] == "verified"


def test_identical_articles_at_a_new_commit_are_unchanged(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = _sources(tmp_path)
    _sync(paths, directory, "a" * 40, DAY_1)
    result = _sync(paths, directory, "b" * 40, DAY_2)
    assert result.status == "unchanged"
    assert result.source_states == {
        "schedule": "unchanged",
        "service_updates": "unchanged",
        "maintenance": "unchanged",
        "pqu_overview": "unchanged",
        "pqu_faq": "unchanged",
    }
    run_health = load_json(paths.run_health_path)
    assert run_health is not None
    assert run_health["checked_commit"] == "b" * 40
    assert run_health["sources"]["service_updates"]["commit"] == "a" * 40


def test_service_update_change_is_recorded_with_its_own_commit(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = _sources(tmp_path)
    _sync(paths, directory, "a" * 40, DAY_1)
    service = directory / SERVICE_FILE
    service.write_text(
        service.read_text(encoding="utf-8").replace(
            "| CY26Q4: 10.0.49\\* | July 27, 2026 | August 17, 2026 | September 11, 2026 |",
            "| CY26Q4: 10.0.49\\* | July 27, 2026 | August 17, 2026 | September 18, 2026 |",
        ),
        encoding="utf-8",
    )
    result = _sync(paths, directory, "e" * 40, DAY_2)
    assert result.source_states == {
        "schedule": "unchanged",
        "service_updates": "updated",
        "maintenance": "unchanged",
        "pqu_overview": "unchanged",
        "pqu_faq": "unchanged",
    }
    changes = load_json(paths.data_dir / "pqu-changes.json")
    assert changes is not None
    service_changes = [c for c in changes["records"] if c["entity"] == "service_update"]
    assert [
        (c["change_type"], c["field"], c["old_value"], c["new_value"], c["source_commit"])
        for c in service_changes
    ] == [
        (
            "modified",
            "10.0.49#general_availability_date",
            "2026-09-11",
            "2026-09-18",
            "e" * 40,
        )
    ]
    # The schedule article did not change, so its provenance stays pinned to the first commit.
    pqu = load_json(paths.pqu_path)
    assert pqu is not None
    assert pqu["source"]["commit"] == "a" * 40
    assert {record["source"]["commit"] for record in pqu["records"]} == {"a" * 40}
    assert verify_artifacts(paths)["status"] == "verified"


def test_verify_rejects_mismatched_source_provenance(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    _sync(paths, _sources(tmp_path), "a" * 40, DAY_1)
    path = paths.data_dir / "service-updates.json"
    document = json.loads(path.read_text(encoding="utf-8"))
    document["source"]["commit"] = "f" * 40
    path.write_text(json.dumps(document), encoding="utf-8")
    with pytest.raises(
        ValidationError, match=r"Source provenance differs in service-updates\.json"
    ):
        verify_artifacts(paths)


def _learn_article(paths: Paths, key: str) -> dict:
    learn = load_json(paths.data_dir / "learn.json")
    assert learn is not None
    return next(article for article in learn["articles"] if article["key"] == key)


def test_faq_and_overview_are_published_as_learn_articles(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    _sync(paths, _sources(tmp_path), "a" * 40, DAY_1)
    faq = _learn_article(paths, "pqu_faq")
    assert faq["title"] == "Proactive quality updates (PQU) - FAQ"
    assert faq["markdown_date"] == "2026-08-10"
    assert faq["file_path"].endswith("get-started/quality-updates-faq.md")
    raw = (LIVE_FIXTURES / "quality-updates-faq.md").read_text(encoding="utf-8")
    questions = [line for line in raw.splitlines() if line.startswith("### ")]
    assert len(questions) == 29
    assert len([s for s in faq["sections"] if s["level"] == 3]) == len(questions)
    overview = _learn_article(paths, "pqu_overview")
    assert [section["id"] for section in overview["sections"]] == [
        "intro",
        "what-are-pqus",
        "why-is-microsoft-introducing-pqus",
        "what-investments-is-microsoft-making-to-enable-safe-deployments-of-pqus",
    ]
    metadata = load_json(paths.metadata_path)
    assert metadata is not None
    assert metadata["sources"]["pqu_faq"]["state"] == "current"
    assert metadata["sources"]["pqu_faq"]["label"] == "Proactive quality updates FAQ"
    assert verify_artifacts(paths)["status"] == "verified"


def test_broken_faq_keeps_its_last_copy_and_reports_it(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = _sources(tmp_path)
    _sync(paths, directory, "a" * 40, DAY_1)
    before = _learn_article(paths, "pqu_faq")
    # Microsoft publishes an unreadable article: front matter only, no title or sections.
    (directory / "quality-updates-faq.md").write_text(
        "---\ntitle: FAQ\nms.date: 09/27/2026\n---\n", encoding="utf-8"
    )
    result = _sync(paths, directory, "b" * 40, DAY_2)
    assert result.source_states["pqu_faq"] == "stale"
    assert result.source_states["pqu_overview"] == "unchanged"
    assert _learn_article(paths, "pqu_faq") == before
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    stale = [item for item in quality["records"] if item["code"] == "source-stale"]
    assert [item["message"].split(" could not")[0] for item in stale] == [
        "Proactive quality updates FAQ"
    ]
    changes = load_json(paths.data_dir / "pqu-changes.json")
    assert changes is not None
    assert [c for c in changes["records"] if c["entity"] == "guidance"] == []
    assert verify_artifacts(paths)["status"] == "verified"


def test_faq_text_change_is_tracked_per_question(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = _sources(tmp_path)
    _sync(paths, directory, "a" * 40, DAY_1)
    faq_path = directory / "quality-updates-faq.md"
    original = faq_path.read_text(encoding="utf-8")
    edited = original.replace(
        "### Can customers request exclusions from the new biweekly cadence?\n\nNo.",
        "### Can customers request exclusions from the new biweekly cadence?\n\nYes, rarely.",
    )
    assert edited != original
    faq_path.write_text(edited, encoding="utf-8")
    _sync(paths, directory, "c" * 40, DAY_2)
    changes = load_json(paths.data_dir / "pqu-changes.json")
    assert changes is not None
    guidance = [c for c in changes["records"] if c["entity"] == "guidance"]
    assert [(c["change_type"], c["field"], c["source_commit"]) for c in guidance] == [
        (
            "modified",
            "pqu_faq#can-customers-request-exclusions-from-the-new-biweekly-cadence",
            "c" * 40,
        )
    ]
    assert guidance[0]["new_value"]["title"] == (
        "Can customers request exclusions from the new biweekly cadence?"
    )
