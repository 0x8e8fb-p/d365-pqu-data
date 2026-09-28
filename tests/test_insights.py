from __future__ import annotations

import csv
import shutil
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from conftest import LIVE_FIXTURES, SCHEMA_DIR
from d365_pqu.config import Paths
from d365_pqu.insights import METRICS, build_insights
from d365_pqu.lifecycle import parse_service_updates
from d365_pqu.normalize import normalize_source
from d365_pqu.parser import parse_source
from d365_pqu.pipeline import load_json, run_sync, verify_artifacts
from d365_pqu.source import bundle_from_directory
from d365_pqu.validation import validate_document

PQU_1_WARNING = (
    "Source warning: Change cutoff 2026-02-04 and train start 2025-02-09 are in different years"
)


@pytest.fixture(scope="module")
def live() -> dict[str, Any]:
    bundle = bundle_from_directory(
        LIVE_FIXTURES, commit="a" * 40, now=datetime(2026, 9, 28, tzinfo=UTC)
    )
    dataset = normalize_source(parse_source(bundle.schedule.markdown), bundle.schedule)
    service_updates, _ = parse_service_updates(bundle.documents["service_updates"].markdown)
    built = build_insights(
        records=dataset.records,
        stations=dataset.stations,
        regions=dataset.regions,
        service_updates=service_updates,
        quality=dataset.quality,
    )
    return {record["id"]: record for record in built["records"]} | {"_built": built}


def _figure(live: dict[str, Any], record_id: str) -> tuple[Any, ...]:
    record = live[record_id]
    return (record["value"], record["sample_size"], record["min"], record["max"])


def test_train_cadence_matches_the_biweekly_change(live) -> None:
    # Microsoft's FAQ: "Starting with version 10.0.47, PQUs move from the previous 28-day
    # schedule to a two-week cadence." The published start dates agree.
    assert _figure(live, "train-cadence:10.0.46") == (28, 6, 28, 35)
    assert live["train-cadence:10.0.46"]["excluded"] == [
        {"item": "10.0.46-PQU-1", "reason": PQU_1_WARNING}
    ]
    assert _figure(live, "train-cadence:10.0.47") == (14, 16, 14, 21)
    assert live["train-cadence:10.0.47"]["excluded"] == []
    assert _figure(live, "train-cadence:10.0.48") == (14, 16, 14, 21)
    assert _figure(live, "train-cadence:10.0.49") == (14, 16, 14, 21)
    assert live["train-cadence:10.0.47"]["summary"] == (
        "10.0.47: a new train starts every 14 days, or 2 weeks "
        "(median of 16 intervals; range 14\u201321 days)."
    )


def test_cutoff_and_train_length(live) -> None:
    assert _figure(live, "cutoff-to-start:10.0.46") == (5, 7, 5, 10)
    assert _figure(live, "cutoff-to-start:10.0.47") == (0, 17, 0, 0)
    assert live["cutoff-to-start:10.0.47"]["summary"] == (
        "10.0.47: the change cutoff falls on the train's start date (median of 17 trains)."
    )
    assert _figure(live, "train-length:10.0.46") == (34, 7, 34, 35)
    assert _figure(live, "train-length:10.0.47") == (25, 17, 25, 25)
    for record_id in ("cutoff-to-start:10.0.46", "train-length:10.0.46"):
        assert [entry["item"] for entry in live[record_id]["excluded"]] == ["10.0.46-PQU-1"]


def test_trains_by_status_follow_microsofts_statuses(live) -> None:
    record = live["trains-by-status:10.0.47"]
    assert record["value"] == 17
    assert record["breakdown"] == {
        "In-Progress": 2,
        "Not Started": 4,
        "Completed": 10,
        "Canceled": 1,
    }
    assert record["summary"] == (
        "10.0.47: 17 trains: 2 In-Progress, 4 Not Started, 10 Completed, 1 Canceled."
    )


def test_station_rollout_order_and_production_gap(live) -> None:
    offsets = {n: live[f"station-offset:station-{n}"]["value"] for n in range(2, 7)}
    assert offsets == {2: 5, 3: 5, 4: 12, 5: 19, 6: 19}
    assert "station-offset:station-1" not in live
    assert _figure(live, "station-offset:station-4") == (12, 5, 12, 14)
    # Microsoft's FAQ: production receives updates at least five days after sandbox.
    assert _figure(live, "sandbox-to-production:all") == (5, 25, 5, 12)
    assert live["sandbox-to-production:all"]["label"] == "All stations"
    assert _figure(live, "sandbox-to-production:station-3") == (5, 5, 5, 11)
    assert "sandbox-to-production:station-1" not in live


def test_regions_per_station_and_opt_in_station(live) -> None:
    counts = {n: live[f"regions-per-station:station-{n}"]["value"] for n in range(2, 7)}
    assert counts == {2: 6, 3: 11, 4: 7, 5: 3, 6: 3}
    metric = next(m for m in live["_built"]["metrics"] if m["id"] == "regions-per-station")
    assert metric["excluded"] == [{"item": "Station 1", "reason": "Only for opted-in environments"}]


def test_service_update_rhythm(live) -> None:
    # Microsoft: four service updates a year; autoupdate windows four weeks apart.
    assert _figure(live, "service-update-interval:all") == (91, 6, 77, 105)
    assert live["service-update-interval:all"]["breakdown"]["10.0.46 to 10.0.47"] == 77
    assert _figure(live, "preview-to-ga:all") == (46, 7, 42, 63)
    assert _figure(live, "ga-to-first-autoupdate:all") == (21, 7, 21, 40)
    assert _figure(live, "first-to-second-autoupdate:all") == (28, 7, 28, 30)
    assert _figure(live, "ga-to-end-of-service:all") == (252, 7, 238, 256)
    assert live["ga-to-end-of-service:all"]["breakdown"]["10.0.46"] == 238
    assert live["ga-to-end-of-service:all"]["summary"] == (
        "End of service comes 252 days, or 36 weeks, after general availability "
        "(median of 7 service updates; range 238\u2013256 days)."
    )


def test_highlights_pick_the_newest_cadence_and_key_rhythms(live) -> None:
    assert live["_built"]["highlights"] == [
        "train-cadence:10.0.49",
        "sandbox-to-production:all",
        "service-update-interval:all",
    ]


def test_records_follow_metric_order_and_carry_units(live) -> None:
    built = live["_built"]
    order = [metric["id"] for metric in METRICS]
    positions = [order.index(record["metric"]) for record in built["records"]]
    assert positions == sorted(positions)
    by_metric = {metric["id"]: metric for metric in built["metrics"]}
    for record in built["records"]:
        assert record["unit"] == by_metric[record["metric"]]["unit"]
        assert record["statistic"] == by_metric[record["metric"]]["statistic"]


def test_empty_data_has_no_figures() -> None:
    built = build_insights(records=[], stations=[], regions=[], service_updates=[], quality=[])
    assert built["records"] == []
    assert built["highlights"] == []
    assert [metric["id"] for metric in built["metrics"]] == [metric["id"] for metric in METRICS]
    assert all(metric["excluded"] == [] for metric in built["metrics"])


def test_single_train_gaps_and_missing_dates_are_explained() -> None:
    def train(number: int, start: str | None, version: str = "10.0.50") -> dict[str, Any]:
        return {
            "pqu_id": f"{version}-PQU-{number}",
            "application_version": version,
            "release_number": number,
            "status": "Not Started",
            "change_cutoff_date": start,
            "train_start_date": start,
            "train_end_date": None,
        }

    records = [
        train(1, "2026-01-07"),
        train(2, "2026-01-21"),
        train(4, "2026-02-18"),
        train(5, None),
        train(1, "2026-03-04", version="10.0.51"),
    ]
    built = build_insights(records=records, stations=[], regions=[], service_updates=[], quality=[])
    by_id = {record["id"]: record for record in built["records"]}
    cadence = by_id["train-cadence:10.0.50"]
    assert (cadence["value"], cadence["sample_size"]) == (14, 1)
    assert (
        cadence["summary"] == "10.0.50: a new train starts every 14 days, or 2 weeks (1 interval)."
    )
    assert cadence["excluded"] == [
        {"item": "10.0.50-PQU-2 to 10.0.50-PQU-4", "reason": "Train numbers are not consecutive"},
        {"item": "10.0.50-PQU-5", "reason": "No published start date"},
    ]
    metric = next(m for m in built["metrics"] if m["id"] == "train-cadence")
    assert metric["excluded"] == [{"item": "10.0.51", "reason": "Only one train is listed"}]
    assert "train-length:10.0.50" not in by_id
    length_metric = next(m for m in built["metrics"] if m["id"] == "train-length")
    assert {"item": "10.0.51-PQU-1", "reason": "No published start or end date"} in (
        length_metric["excluded"]
    )


def _paths(tmp_path: Path) -> Paths:
    return Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        excel_dir=tmp_path / "excel",
        schema_dir=SCHEMA_DIR,
        site_dir=tmp_path / "build" / "site",
        tests_dir=tmp_path / "tests",
    )


def test_pipeline_publishes_insights_in_json_csv_and_excel(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = tmp_path / "sources"
    shutil.copytree(LIVE_FIXTURES, directory)
    now = datetime(2026, 9, 28, 3, 0, tzinfo=UTC)
    run_sync(paths, bundle=bundle_from_directory(directory, commit="a" * 40, now=now), now=now)
    document = load_json(paths.data_dir / "insights.json")
    assert document is not None
    validate_document(document, SCHEMA_DIR / "insights.schema.json")
    assert document["count"] == len(document["records"]) == 37
    assert document["source"] == load_json(paths.pqu_path)["source"]  # type: ignore[index]
    assert [category["id"] for category in document["categories"]] == [
        "trains",
        "stations",
        "service-updates",
    ]
    with (paths.data_dir / "insights.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 37
    cadence = next(row for row in rows if row["id"] == "train-cadence:10.0.46")
    assert cadence["value"] == "28"
    assert "10.0.46-PQU-1" in cadence["excluded"]
    workbook = load_workbook(paths.workbook_path, read_only=True)
    sheet = workbook["16_INSIGHTS"]
    assert sheet.max_row == 38
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header[:3] == ["Figure", "Group", "Value"]
    workbook.close()
    assert verify_artifacts(paths)["status"] == "verified"

    # Figures depend only on published data: a later run at a new commit (inside the monthly
    # heartbeat window) changes nothing, even though "today" has moved on three weeks.
    later = datetime(2026, 10, 19, 3, 0, tzinfo=UTC)
    result = run_sync(
        paths, bundle=bundle_from_directory(directory, commit="b" * 40, now=later), now=later
    )
    assert result.status == "unchanged"
    assert load_json(paths.data_dir / "insights.json") == document
