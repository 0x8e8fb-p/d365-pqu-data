from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from conftest import FIXTURES, SCHEMA_DIR
from d365_pqu.config import Paths
from d365_pqu.pipeline import EXPECTED_SHEETS, run_sync


def _paths(tmp_path: Path) -> Paths:
    return Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        excel_dir=tmp_path / "excel",
        schema_dir=SCHEMA_DIR,
        site_dir=tmp_path / "build" / "site",
        tests_dir=tmp_path / "tests",
    )


def _sync(tmp_path: Path) -> Paths:
    import hashlib

    from d365_pqu.models import SourceDocument
    from d365_pqu.source import parse_markdown_date

    paths = _paths(tmp_path)
    markdown = (FIXTURES / "source-minimal.md").read_text(encoding="utf-8")
    document = SourceDocument(
        markdown=markdown,
        source_commit="a" * 40,
        article_url="https://learn.microsoft.com/example",
        raw_url="https://raw.githubusercontent.com/example/a/schedule.md",
        markdown_date=parse_markdown_date(markdown),
        retrieved_at=datetime(2026, 9, 24, 12, 0, tzinfo=UTC),
        sha256=hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
    )
    run_sync(paths, source=document, now=datetime(2026, 9, 24, 12, 0, tzinfo=UTC))
    return paths


def test_workbook_contains_expected_sheets(tmp_path: Path) -> None:
    paths = _sync(tmp_path)
    wb = load_workbook(paths.workbook_path)
    assert tuple(wb.sheetnames) == EXPECTED_SHEETS
    assert wb["00_README"]["A1"].value == "D365 Finance & Operations PQU Tracker"
    assert wb["01_PQU_MASTER"].max_row == 5
    assert wb["02_CURRENT_PQU"].max_row == 3
    assert wb["03_UPCOMING_PQU"].max_row == 2
    assert wb["04_STATION_SCHEDULE"].max_row == 13
    assert wb["05_REGION_MAPPING"].max_row == 10
    wb.close()


def test_workbook_cells_are_static_and_typed(tmp_path: Path) -> None:
    paths = _sync(tmp_path)
    wb = load_workbook(paths.workbook_path)
    ws = wb["01_PQU_MASTER"]
    header = [cell.value for cell in ws[1]]
    row = {header[index]: cell for index, cell in enumerate(ws[2])}
    assert row["PQU ID"].value in {"10.0.47-PQU-2", "10.0.48-PQU-5"}
    assert row["Status"].fill.fgColor.rgb
    for sheet in wb.worksheets:
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                if isinstance(cell.value, str):
                    assert not cell.value.startswith("=")
    wb.close()


def test_workbook_records_source_and_quality(tmp_path: Path) -> None:
    paths = _sync(tmp_path)
    wb = load_workbook(paths.workbook_path)
    metadata_values = [cell.value for row in wb["11_SOURCE_METADATA"].iter_rows() for cell in row]
    assert "a" * 40 in metadata_values
    assert "Data Quality" not in metadata_values
    wb.close()
