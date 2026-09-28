from __future__ import annotations

import csv
import shutil
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from conftest import LIVE_FIXTURES, SCHEMA_DIR
from d365_pqu.config import Paths
from d365_pqu.events import build_events
from d365_pqu.lifecycle import parse_service_updates
from d365_pqu.normalize import normalize_source
from d365_pqu.parser import parse_source
from d365_pqu.pipeline import load_json, run_sync, verify_artifacts
from d365_pqu.source import bundle_from_directory
from d365_pqu.validation import validate_document

SCHEDULE_URL = (
    "https://learn.microsoft.com/en-us/dynamics365/fin-ops-core/dev-itpro/get-started/"
    "quality-updates-schedule"
)


def _live_events(markdown: str | None = None) -> list[dict[str, Any]]:
    from d365_pqu.learn import extract_article
    from d365_pqu.lifecycle import SERVICE_UPDATE_RULES

    bundle = bundle_from_directory(
        LIVE_FIXTURES, commit="a" * 40, now=datetime(2026, 9, 28, tzinfo=UTC)
    )
    schedule = bundle.schedule
    text = markdown or schedule.markdown
    dataset = normalize_source(parse_source(text), schedule)
    service = bundle.documents["service_updates"]
    service_updates, _ = parse_service_updates(service.markdown)
    service_article = extract_article(
        service.markdown,
        key=service.key,
        file_path=service.file_path,
        article_url=service.article_url,
        rules=SERVICE_UPDATE_RULES,
    )
    return build_events(
        records=dataset.records,
        stations=dataset.stations,
        service_updates=service_updates,
        articles=[*dataset.articles, service_article],
        quality=dataset.quality,
    )


@pytest.fixture(scope="module")
def live() -> list[dict[str, Any]]:
    return _live_events()


def test_every_published_date_becomes_one_event(live) -> None:
    kinds = Counter(event["kind"] for event in live)
    assert kinds == {
        "change_cutoff": 59,
        "train_window": 59,
        "sandbox_window": 30,
        "production_window": 25,
        "preview": 7,
        "preview_latest_update": 7,
        "general_availability": 7,
        "autoupdate_first": 7,
        "autoupdate_second": 7,
        "end_of_service": 7,
    }
    assert len({event["id"] for event in live}) == len(live) == 215
    starts = [event["start_date"] for event in live]
    assert starts == sorted(starts)


def test_event_shapes_and_links(live) -> None:
    by_id = {event["id"]: event for event in live}
    window = by_id["production_window:10.0.48-PQU-6:station-4"]
    assert window == {
        "id": "production_window:10.0.48-PQU-6:station-4",
        "kind": "production_window",
        "category": "station",
        "title": "10.0.48 PQU-6 · Station 4 production",
        "start_date": "2026-10-03",
        "end_date": "2026-10-04",
        "pqu_id": "10.0.48-PQU-6",
        "application_version": "10.0.48",
        "station": 4,
        "status": "In-Progress",
        "source_key": "schedule",
        "url": f"{SCHEDULE_URL}#schedule",
        "warnings": [],
    }
    cutoff = by_id["change_cutoff:10.0.48-PQU-7"]
    assert (cutoff["start_date"], cutoff["end_date"], cutoff["title"]) == (
        "2026-09-30",
        "2026-09-30",
        "10.0.48 PQU-7 change cutoff",
    )
    assert cutoff["url"] == f"{SCHEDULE_URL}#high-level-pqu-train-schedule"
    eos = by_id["end_of_service:10.0.46"]
    assert (eos["start_date"], eos["category"], eos["pqu_id"], eos["status"]) == (
        "2026-08-21",
        "service_update",
        None,
        None,
    )
    assert eos["url"].endswith(
        "public-preview-releases#targeted-release-schedule-dates-subject-to-change"
    )
    without_article = build_events(
        records=[],
        stations=[],
        service_updates=[{"version": "10.0.46", **{"end_of_service_date": "2026-08-21"}}],
        articles=[],
        quality=[],
    )
    assert without_article[0]["url"] == (
        "https://learn.microsoft.com/en-us/dynamics365/fin-ops-core/dev-itpro/get-started/"
        "public-preview-releases"
    )
    # Station 1 has no production window in Microsoft's schedule.
    assert "production_window:10.0.48-PQU-6:station-1" not in by_id
    assert "sandbox_window:10.0.48-PQU-6:station-1" in by_id


def test_source_warnings_travel_with_the_affected_events(live) -> None:
    flagged = {event["id"]: event["warnings"] for event in live if event["warnings"]}
    assert flagged == {
        "change_cutoff:10.0.46-PQU-1": ["cutoff-start-year-mismatch"],
        "train_window:10.0.46-PQU-1": ["cutoff-start-year-mismatch"],
    }


def test_ids_stay_the_same_when_microsoft_moves_a_date(live) -> None:
    markdown = (LIVE_FIXTURES / "quality-updates-schedule.md").read_text(encoding="utf-8")
    moved = markdown.replace(
        "| Station 4 | September 28 to October 1, 2026 | October 3 to October 4, 2026 |",
        "| Station 4 | September 29 to October 2, 2026 | October 10 to October 11, 2026 |",
    )
    assert moved != markdown
    after = {event["id"]: event for event in _live_events(moved)}
    before = {event["id"]: event for event in live}
    assert set(after) == set(before)
    changed = sorted(
        event_id
        for event_id in before
        if before[event_id]["start_date"] != after[event_id]["start_date"]
    )
    assert changed == [
        "production_window:10.0.48-PQU-6:station-4",
        "sandbox_window:10.0.48-PQU-6:station-4",
    ]
    assert after["production_window:10.0.48-PQU-6:station-4"]["start_date"] == "2026-10-10"


def test_no_published_data_means_no_events() -> None:
    assert build_events(records=[], stations=[], service_updates=[], articles=[], quality=[]) == []


def _paths(tmp_path: Path) -> Paths:
    return Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        excel_dir=tmp_path / "excel",
        schema_dir=SCHEMA_DIR,
        site_dir=tmp_path / "build" / "site",
        tests_dir=tmp_path / "tests",
    )


def test_pipeline_publishes_key_dates(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = tmp_path / "sources"
    shutil.copytree(LIVE_FIXTURES, directory)
    now = datetime(2026, 9, 28, 3, 0, tzinfo=UTC)
    run_sync(paths, bundle=bundle_from_directory(directory, commit="a" * 40, now=now), now=now)
    document = load_json(paths.data_dir / "events.json")
    assert document is not None
    validate_document(document, SCHEMA_DIR / "event.schema.json")
    assert document["count"] == 215
    with (paths.data_dir / "events.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 215
    row = next(item for item in rows if item["id"] == "production_window:10.0.48-PQU-6:station-4")
    assert (row["start_date"], row["station"], row["warnings"]) == ("2026-10-03", "4", "[]")
    workbook = load_workbook(paths.workbook_path, read_only=True)
    sheet = workbook["17_KEY_DATES"]
    assert sheet.max_row == 216
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header[:3] == ["Start Date", "End Date", "Event"]
    workbook.close()
    assert verify_artifacts(paths)["status"] == "verified"
