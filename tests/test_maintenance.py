from __future__ import annotations

import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import LIVE_FIXTURES, SCHEMA_DIR
from d365_pqu.config import Paths
from d365_pqu.errors import ParserError
from d365_pqu.maintenance import (
    assign_maintenance_geos,
    match_region,
    parse_days,
    parse_duration_hours,
    parse_maintenance_windows,
    parse_start_time,
)
from d365_pqu.models import RegionRecord
from d365_pqu.pipeline import load_json, run_sync, verify_artifacts
from d365_pqu.source import bundle_from_directory

LIVE = (LIVE_FIXTURES / "plannedmaintenance-selfservice.md").read_text(encoding="utf-8")
GEOS = [
    "Australia",
    "Asia",
    "Brazil",
    "Canada",
    "China",
    "Europe",
    "France",
    "India",
    "Japan",
    "Norway",
    "South Africa",
    "Switzerland",
    "United Arab Emirates",
    "United Kingdom",
    "United States",
]
# Every region in Microsoft's station-to-region mapping (live snapshot) and its geography.
EXPECTED_GEOS = {
    "Canada Central": "Canada",
    "Canada East": "Canada",
    "France Central": "France",
    "India Central": "India",
    "Norway East": "Norway",
    "Switzerland West": "Switzerland",
    "France South": "France",
    "India South": "India",
    "Norway West": "Norway",
    "Switzerland North": "Switzerland",
    "South Africa North": "South Africa",
    "Australia East": "Australia",
    "UK South": "United Kingdom",
    "UAE North": "United Arab Emirates",
    "Japan East": "Japan",
    "Australia South East": "Australia",
    "South East Asia": "Asia",
    "East Asia": "Asia",
    "UK West": "United Kingdom",
    "Japan West": "Japan",
    "Brazil South": "Brazil",
    "North Europe": "Europe",
    "East US": "United States",
    "UAE Central": "United Arab Emirates",
    "West Europe": "Europe",
    "Central US": "United States",
    "West US": "United States",
    "China": "China",
    "DoD": None,
    "Government Community Cloud": None,
}


def _region(name: str, station: int = 2, is_region: bool = True) -> RegionRecord:
    return RegionRecord(
        station=station,
        station_label=f"Station {station}",
        region=name,
        is_region=is_region,
        maintenance_geo=None,
    )


def test_live_article_lists_fifteen_geographies_in_microsofts_order() -> None:
    records, quality = parse_maintenance_windows(LIVE)
    assert quality == []
    assert [record["geo"] for record in records] == GEOS
    by_geo = {record["geo"]: record for record in records}
    assert by_geo["Europe"] == {
        "geo": "Europe",
        "start_time_utc": "22:00",
        "days": ["Friday", "Saturday"],
        "duration_hours": 6,
        "duration_text": "Six hours",
    }
    assert by_geo["India"]["start_time_utc"] == "18:30"
    assert by_geo["United States"]["days"] == ["Saturday", "Sunday"]
    assert {record["duration_hours"] for record in records} == {6}


@pytest.mark.parametrize(
    ("text", "hours"),
    [
        ("Six hours", 6),
        ("six Hours", 6),
        ("6 hours", 6),
        ("6h", 6),
        ("1 hour", 1),
        ("Twenty-four hours", 24),
        ("90 minutes", 1.5),
        ("2.5 hrs", 2.5),
        ("Several hours", None),
        ("25 hours", None),
        ("0 hours", None),
        ("", None),
    ],
)
def test_duration_parsing(text: str, hours: float | None) -> None:
    assert parse_duration_hours(text) == hours


def test_start_time_and_day_parsing() -> None:
    assert parse_start_time("13:00 UTC") == "13:00"
    assert parse_start_time("4:00 utc") == "04:00"
    assert parse_start_time("18:30") == "18:30"
    assert parse_start_time("22:00 PST") is None
    assert parse_start_time("24:00 UTC") is None
    assert parse_days("Friday, Saturday") == ["Friday", "Saturday"]
    assert parse_days("Sat and Sun") == ["Saturday", "Sunday"]
    assert parse_days("Fri / Sat; Fri") == ["Friday", "Saturday"]
    assert parse_days("Weekends") is None
    assert parse_days("") is None


def test_unreadable_rows_are_warnings_and_a_missing_table_is_an_error() -> None:
    edited = (
        LIVE.replace("| Brazil | 04:00 UTC |", "| Brazil | early morning |")
        .replace("| Canada | 04:00 UTC | Saturday, Sunday |", "| Canada | 04:00 UTC | Weekends |")
        .replace(
            "| Japan | 16:00 UTC | Friday, Saturday | Six hours |",
            "| Japan | 16:00 UTC | Friday, Saturday | A while |",
        )
        .replace("| Norway |", "| Europe |")
    )
    records, quality = parse_maintenance_windows(edited)
    geos = [record["geo"] for record in records]
    assert "Brazil" not in geos and "Canada" not in geos
    assert next(r for r in records if r["geo"] == "Japan")["duration_hours"] is None
    assert [(item["code"], item["field"]) for item in quality] == [
        ("maintenance-window-row-invalid", "Brazil"),
        ("maintenance-window-row-invalid", "Canada"),
        ("maintenance-window-duration-invalid", "Japan#duration_hours"),
        ("maintenance-window-duplicate", "Europe"),
    ]
    with pytest.raises(ParserError, match="not found or its headers changed"):
        parse_maintenance_windows(LIVE.replace("| Geo |", "| Geography |"))


def test_every_live_region_maps_to_the_expected_geography() -> None:
    regions = json.loads((LIVE_FIXTURES / "expected-regions.json").read_text(encoding="utf-8"))
    names = [row["region"] for row in regions if row["is_region"]]
    assert sorted(names) == sorted(EXPECTED_GEOS)
    assert {name: match_region(name, GEOS).geo for name in names} == EXPECTED_GEOS


def test_sovereign_and_unknown_regions_are_reported_not_guessed() -> None:
    regions = [
        _region("DoD", 6),
        _region("Government Community Cloud", 6),
        _region("Mars Central"),
        _region("Only for opted-in environments", 1, is_region=False),
        _region("North Europe"),
    ]
    windows = [{"geo": geo} for geo in GEOS]
    quality = assign_maintenance_geos(regions, windows)
    assert [region["maintenance_geo"] for region in regions] == [None, None, None, None, "Europe"]
    assert [(item["code"], item["severity"], item["field"]) for item in quality] == [
        ("maintenance-geo-unmapped", "info", "region.DoD"),
        ("maintenance-geo-unmapped", "info", "region.Government Community Cloud"),
        ("maintenance-geo-unmatched", "warning", "region.Mars Central"),
    ]
    assert quality[0]["message"] == (
        "Microsoft's planned maintenance window table does not list a geography for DoD, "
        "so no maintenance window is shown for it."
    )


def test_matching_prefers_the_longest_name_and_flags_ties() -> None:
    assert match_region("South Africa North", [*GEOS, "Africa"]).geo == "South Africa"
    tie = match_region("Europe France Central", GEOS)
    assert (tie.geo, tie.reason, tie.candidates) == (None, "ambiguous", ("Europe", "France"))
    # A government region is never given a commercial geography's window ...
    assert match_region("US Gov Virginia", GEOS).reason == "sovereign"
    # ... but maps when Microsoft lists a matching government geography.
    assert match_region("DoD", [*GEOS, "DoD"]).geo == "DoD"
    assert match_region("US Gov Virginia", [*GEOS, "US Government"]).geo == "US Government"
    assert match_region("Korea Central", GEOS).reason == "unmatched"


def _paths(tmp_path: Path) -> Paths:
    return Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        excel_dir=tmp_path / "excel",
        schema_dir=SCHEMA_DIR,
        site_dir=tmp_path / "build" / "site",
        tests_dir=tmp_path / "tests",
    )


def _sync(paths: Paths, directory: Path, commit: str, day: int):
    now = datetime(2026, 9, day, 12, 0, tzinfo=UTC)
    return run_sync(paths, bundle=bundle_from_directory(directory, commit=commit, now=now), now=now)


def test_pipeline_publishes_windows_and_region_geographies(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    directory = tmp_path / "sources"
    shutil.copytree(LIVE_FIXTURES, directory)
    _sync(paths, directory, "a" * 40, 24)
    document = load_json(paths.data_dir / "maintenance-windows.json")
    assert document is not None
    assert (document["state"], document["count"]) == ("current", 15)
    assert document["source"]["file_path"].endswith("deployment/plannedmaintenance-selfservice.md")
    assert document["source"]["markdown_date"] == "2026-04-02"
    regions = load_json(paths.data_dir / "pqu-regions.json")
    assert regions is not None
    mapped = {
        row["region"]: row["maintenance_geo"] for row in regions["records"] if row["is_region"]
    }
    assert mapped == EXPECTED_GEOS
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    assert quality["warning_count"] == 1
    assert [item["code"] for item in quality["records"] if item["severity"] == "info"] == [
        "maintenance-geo-unmapped",
        "maintenance-geo-unmapped",
    ]
    health = load_json(paths.health_path)
    assert health is not None
    assert health["warning_count"] == 1
    learn = load_json(paths.data_dir / "learn.json")
    assert learn is not None
    article = next(item for item in learn["articles"] if item["key"] == "maintenance")
    windows = next(s for s in article["sections"] if s["aliases"] == ["windows"])
    assert {"type": "dataset", "name": "maintenance_windows"} in windows["blocks"]
    assert verify_artifacts(paths)["status"] == "verified"
    workbook = load_workbook(paths.workbook_path, read_only=True)
    assert workbook["15_MAINTENANCE_WINDOWS"].max_row == 16
    header = [cell.value for cell in next(workbook["05_REGION_MAPPING"].iter_rows(max_row=1))]
    assert header[-1] == "Maintenance Geo"
    workbook.close()
    csv_header = (paths.data_dir / "maintenance-windows.csv").read_text(encoding="utf-8")
    assert csv_header.splitlines()[0] == (
        "geo,start_time_utc,days,duration_hours,duration_text,source_commit,source_url"
    )

    # Microsoft moves the Europe window: a maintenance_window change, regions unchanged.
    article_path = directory / "plannedmaintenance-selfservice.md"
    article_path.write_text(
        article_path.read_text(encoding="utf-8").replace(
            "| Europe | 22:00 UTC |", "| Europe | 21:00 UTC |"
        ),
        encoding="utf-8",
    )
    result = _sync(paths, directory, "b" * 40, 25)
    assert result.source_states["maintenance"] == "updated"
    changes = load_json(paths.data_dir / "pqu-changes.json")
    assert changes is not None
    window_changes = [c for c in changes["records"] if c["entity"] == "maintenance_window"]
    assert [
        (c["change_type"], c["field"], c["old_value"], c["new_value"], c["source_commit"])
        for c in window_changes
    ] == [("modified", "Europe#start_time_utc", "22:00", "21:00", "b" * 40)]
    assert [c for c in changes["records"] if c["entity"] == "region"] == []

    # Microsoft drops the India row: the window is removed and the two India regions change.
    article_path.write_text(
        article_path.read_text(encoding="utf-8").replace(
            "| India | 18:30 UTC | Friday, Saturday | Six hours |\n", ""
        ),
        encoding="utf-8",
    )
    _sync(paths, directory, "c" * 40, 26)
    changes = load_json(paths.data_dir / "pqu-changes.json")
    assert changes is not None
    latest = [c for c in changes["records"] if c["source_commit"] == "c" * 40]
    assert sorted((c["entity"], c["change_type"], c["field"]) for c in latest) == [
        ("maintenance_window", "removed", "India"),
    ]
    region_changes = [
        c
        for c in changes["records"]
        if c["entity"] == "region" and c["changed_at"].startswith("2026-09-26")
    ]
    assert sorted(c["field"] for c in region_changes) == [
        "region.India Central",
        "region.India South",
    ]
    assert all(c["new_value"]["maintenance_geo"] is None for c in region_changes)
    quality = load_json(paths.data_dir / "quality-report.json")
    assert quality is not None
    unmatched = [item for item in quality["records"] if item["code"] == "maintenance-geo-unmatched"]
    assert [item["field"] for item in unmatched] == ["region.India Central", "region.India South"]
    assert verify_artifacts(paths)["status"] == "verified"
